"""Validate the revised BMC Genomics submission package without altering analysis files."""
from __future__ import annotations
import json
import re
import sys
import zipfile
from pathlib import Path
from PIL import Image
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission_bmc_genomics_revised"
sys.path.insert(0, str(ROOT / "manuscript"))
from manuscript_content import ABSTRACT, FIGURE_LEGENDS, REFERENCES, TITLE


def doc_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])


main = doc_text(OUT / "01_Main_Manuscript_BMC_Genomics.docx")
title_page = doc_text(OUT / "02_Title_Page.docx")
cover = doc_text(OUT / "03_Cover_Letter_BMC_Genomics.docx")
sheet = doc_text(OUT / "04_Submission_Copy_Paste_Sheet.docx")
supp = doc_text(OUT / "additional_files/Additional_file_1_Supplementary_Materials.docx")
body = main.split("References\n")[0]
checks = {}
checks["title_synchronized"] = all(TITLE in x for x in (main, title_page, cover, sheet, supp))
checks["new_author_order_present"] = "Yaobin He, Youxing Huang, Rong Chen, Wei He, Kairui Liu and Yipei Huang" in main
checks["corresponding_author_updated"] = all("huangyp2025@163.com" in x for x in (main, title_page, cover, sheet))
checks["old_contact_absent"] = "waiqike7" not in main + title_page + cover + sheet
checks["funding_updated"] = all("20252011" in x and "20251174" in x and "20264020" in x for x in (main, sheet))
checks["structured_abstract_synchronized"] = all(value in main and value in sheet for value in ABSTRACT.values())
checks["abstract_under_350_words"] = sum(len(v.split()) for v in ABSTRACT.values()) <= 350

cited = set()
for group in re.findall(r"\[([\d,\s-]+)\]", body):
    for item in group.split(","):
        item = item.strip()
        if "-" in item:
            a, b = map(int, item.split("-")); cited.update(range(a, b + 1))
        elif item:
            cited.add(int(item))
checks["all_references_cited"] = cited == set(range(1, len(REFERENCES) + 1))
positions = [body.find(f"Figure {i}") for i in range(1, 7)]
checks["figures_first_cited_in_order"] = all(x >= 0 for x in positions) and positions == sorted(positions)
checks["six_legends"] = len(FIGURE_LEGENDS) == 6

expected = ["Figure1_study_design", "Figure2_bulk_outcomes", "Figure3_healthy_reference",
            "Figure4_composition_vs_state", "Figure5_kitagawa_decomposition", "Figure6_reference_programs",
            "Supplementary_Figure_S1_threshold_sensitivity", "Supplementary_Figure_S2_DecontX",
            "Supplementary_Figure_S3_ambient_controls"]
checks["figure_inventory_complete"] = all((OUT / "figures" / f"{stem}.{ext}").exists() for stem in expected for ext in ("png", "pdf"))
dpi_ok = True
for stem in expected:
    with Image.open(OUT / "figures" / f"{stem}.png") as image:
        dpi = image.info.get("dpi", (0, 0))
        dpi_ok &= min(dpi) >= 299
checks["png_300_dpi"] = bool(dpi_ok)
checks["primary_validation_result_retained"] = "fixed support criterion was not met" in main
checks["no_reviewer_response_language"] = not re.search(r"reviewer-requested|in response to the reviewer", main + supp, re.I)
with zipfile.ZipFile(OUT / "01_Main_Manuscript_BMC_Genomics.docx") as archive:
    document_xml = archive.read("word/document.xml").decode("utf-8")
checks["no_forced_page_breaks_in_main"] = 'w:type="page"' not in document_xml
checks["analysis_plan_present_and_valid_json"] = isinstance(json.loads((ROOT / "planning/analysis_plan_005_health_function.json").read_text(encoding="utf-8")), dict)
checks["source_data_present"] = len(list((OUT / "source_data").glob("*.tsv"))) >= 12
manifest = json.loads((OUT / "PACKAGE_MANIFEST.json").read_text(encoding="utf-8"))
checks["manifest_counts"] = manifest["main_figures"] == 6 and manifest["supplementary_figures"] == 3 and manifest["main_tables"] == 2

result = {"passed": all(checks.values()), "n_checks": len(checks), "checks": checks,
          "failed": [key for key, value in checks.items() if not value]}
(OUT / "quality_checks").mkdir(exist_ok=True)
(OUT / "quality_checks/PACKAGE_VALIDATION.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, indent=2))
if not result["passed"]: raise SystemExit(1)
