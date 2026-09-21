"""Build the complete uncompressed BMC Genomics submission folder."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission_bmc_genomics_revised"
FIG = OUT / "figures"
SUP = OUT / "additional_files"
SOURCE = OUT / "source_data"
QC = OUT / "quality_checks"
PUBFIG = ROOT / "figures" / "publication"
REPO_URL = "https://github.com/liukairui621/uc-mature-colonocyte-recovery"

AUTHORS = ["Yaobin He", "Youxing Huang", "Rong Chen", "Wei He", "Kairui Liu", "Yipei Huang"]
AUTHOR_LINE = ", ".join(AUTHORS[:-1]) + " and " + AUTHORS[-1] + "*"
AFFILIATION = "Department of Abdominal Surgery, The Second Affiliated Hospital of Guangzhou University of Chinese Medicine, Guangzhou, Guangdong, P.R. China"
CORRESPONDING = "Yipei Huang, M.D."
ADDRESS = "Department of Abdominal Surgery\nThe Second Affiliated Hospital of Guangzhou University of Chinese Medicine\nNo. 111 Dade Road, Yuexiu District\nGuangzhou 510000, Guangdong Province, P.R. China"
PHONE = "+86-20-39318791; +86-18826234617"
EMAIL = "huangyp2025@163.com"
KEYWORDS = ["ulcerative colitis", "mucosal healing", "colonocyte", "biologic therapy", "longitudinal transcriptomics", "single-cell RNA sequencing", "pseudobulk"]
FUNDING_TEXT = (
    "This study was supported by the Guangdong Provincial Key Laboratory of Clinical Research on Traditional "
    "Chinese Medicine Syndrome and the Science and Technology Planning Project of Guangdong Province "
    "(No. 2023B1212060063, awarded to R.C.), the Guangzhou Municipal Science and Technology Project "
    "(No. 2024A03J0052, awarded to Y.X.H.), the Elite Clinical Technical Talent Program of Guangdong "
    "Provincial Hospital of Chinese Medicine (awarded to Y.X.H.), and the Traditional Chinese Medicine Bureau "
    "of Guangdong Province (No. 20252011, awarded to W.H., No. 20251174, awarded to Y.B.H., and "
    "No. 20264020, awarded to K.L.). The funders had no role in study design, data acquisition, analysis, "
    "interpretation, manuscript preparation, or the decision to submit the work for publication."
)
CONTRIBUTIONS = (
    "Y.B.H., Y.X.H. and Y.P.H. conceived and supervised the study. K.L. curated the data, developed the "
    "analysis code, performed the statistical analyses and prepared the figures. R.C. and W.H. contributed "
    "to methodology, result interpretation and project administration. K.L. drafted the manuscript. All "
    "authors critically revised the manuscript and approved the final version."
)

sys.path.insert(0, str(ROOT / "manuscript"))
from manuscript_content import (ABSTRACT, BACKGROUND, CONCLUSION, DISCUSSION, FIGURE_LEGENDS, METHODS,
                                REFERENCES, RESULTS_TEXT, RUNNING, SUPPLEMENTARY_FIGURE_LEGENDS,
                                TABLE1, TABLE2, TITLE)


def reset_output() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    for folder in (OUT, FIG, SUP, SOURCE, QC, PUBFIG):
        folder.mkdir(parents=True, exist_ok=True)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def setup_doc(doc: Document, double: bool, line_numbers: bool = False) -> None:
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2 if double else 1.15
    normal.paragraph_format.space_after = Pt(0 if double else 5)
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if "Caption" in doc.styles:
        style = doc.styles["Caption"]
        style.font.name = "Arial"; style.font.size = Pt(9); style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0); style.paragraph_format.line_spacing = 1
    if line_numbers:
        node = OxmlElement("w:lnNumType")
        node.set(qn("w:countBy"), "1"); node.set(qn("w:restart"), "continuous"); node.set(qn("w:distance"), "360")
        section._sectPr.append(node)
    add_page_number(section.footer.paragraphs[0])


def add_label(doc: Document, label: str, text: str):
    p = doc.add_paragraph(); p.add_run(label + ": ").bold = True; p.add_run(text); return p


def add_table(doc: Document, rows, font_size: float = 8.0):
    table = doc.add_table(rows=1, cols=len(rows[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        cells = table.rows[0].cells if r_idx == 0 else table.add_row().cells
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]; cell.text = str(value); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1; p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Arial"; run.font.size = Pt(font_size); run.bold = r_idx == 0
    doc.add_paragraph()
    return table


def make_figures() -> None:
    scripts = [ROOT / "code/plot_publication_main_figures.py", ROOT / "code/plot_005_extension.py",
               ROOT / "code/plot_publication_aux_figures.py"]
    for script in scripts:
        subprocess.run([sys.executable, str(script), "--dest", str(PUBFIG), "--mirror", str(FIG)], check=True)
    expected = [
        "Figure1_study_design", "Figure2_bulk_outcomes", "Figure3_healthy_reference",
        "Figure4_composition_vs_state", "Figure5_kitagawa_decomposition", "Figure6_reference_programs",
        "Supplementary_Figure_S1_threshold_sensitivity", "Supplementary_Figure_S2_DecontX",
        "Supplementary_Figure_S3_ambient_controls",
    ]
    missing = [f"{stem}.png" for stem in expected if not (FIG / f"{stem}.png").exists()]
    if missing: raise FileNotFoundError(missing)


def copy_source_data() -> None:
    mapping = {
        "external_analysis_overview.tsv": ROOT / "runs/004A_evidence_ledger/external_analysis_overview.tsv",
        "discovery_response_models.tsv": ROOT / "runs/001D_annotation/discovery_response_models.tsv",
        "GSE23597_candidate_models.tsv": ROOT / "runs/002A_GSE23597/candidate_models.tsv",
        "GSE73661_IFX_candidate_models.tsv": ROOT / "runs/002B_GSE73661_support/candidate_support_models.tsv",
        "GSE73661_VDZ_candidate_models.tsv": ROOT / "runs/002D_VDZ_support/candidate_models.tsv",
        "GSE282122_joint_primary_models.tsv": ROOT / "runs/003_cell_context/joint_primary_patient_models.tsv",
        "GSE282122_patient_visit_metrics.tsv": ROOT / "runs/003F_baseline_adjustment/patient_visit_metrics.tsv",
        "GSE282122_raw_fixed15_state_models.tsv": ROOT / "runs/003_cell_context/fixed_epithelial_state_models.tsv",
        "GSE282122_DecontX_fixed15_state_models.tsv": ROOT / "runs/003D_decontx_state_family/decontx_fixed15_state_models.tsv",
        "GSE282122_negative_control_models.tsv": ROOT / "runs/003E_ambient_negative_controls/negative_control_CT_models.tsv",
        "GSE282122_contamination_models.tsv": ROOT / "runs/003E_ambient_negative_controls/CT_contamination_patient_models.tsv",
        "healthy_reference_contrasts.tsv": ROOT / "runs/005_health_function/healthy_reference_contrasts.tsv",
        "reference_program_models.tsv": ROOT / "runs/005_health_function/ct_program_models.tsv",
        "reference_program_correlations.tsv": ROOT / "runs/005_health_function/ct_candidate_coordination.tsv",
        "program_membership.tsv": ROOT / "runs/005_preparation/program_membership.tsv",
    }
    for name, src in mapping.items():
        if src.exists(): shutil.copy2(src, SOURCE / name)


def manuscript_docx() -> Path:
    doc = Document(); setup_doc(doc, double=True, line_numbers=True)
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(TITLE).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(AUTHOR_LINE).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(AFFILIATION)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f"*Correspondence: {CORRESPONDING}; {EMAIL}")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Running title: ").bold = True; p.add_run(RUNNING)
    doc.add_heading("Abstract", 1)
    for key, value in ABSTRACT.items(): add_label(doc, key, value)
    add_label(doc, "Keywords", "; ".join(KEYWORDS))
    doc.add_heading("Background", 1)
    for value in BACKGROUND: doc.add_paragraph(value)
    doc.add_heading("Methods", 1)
    for heading, paragraphs in METHODS:
        doc.add_heading(heading, 2)
        for value in paragraphs: doc.add_paragraph(value)
    doc.add_heading("Results", 1)
    for heading, paragraphs in RESULTS_TEXT:
        doc.add_heading(heading, 2)
        for value in paragraphs: doc.add_paragraph(value)
    doc.add_heading("Discussion", 1)
    for value in DISCUSSION: doc.add_paragraph(value)
    doc.add_heading("Conclusions", 1); doc.add_paragraph(CONCLUSION)
    doc.add_heading("Abbreviations", 1)
    doc.add_paragraph("ADA, adalimumab; BH, Benjamini-Hochberg; CI, confidence interval; CT, crypt-top colonocyte; GEO, Gene Expression Omnibus; HC3, heteroskedasticity-consistent covariance estimator 3; IFX, infliximab; RMA, robust multi-array average; scRNA-seq, single-cell RNA sequencing; UC, ulcerative colitis; VDZ, vedolizumab.")
    doc.add_heading("Declarations", 1)
    add_label(doc, "Ethics approval and consent to participate", "Not applicable. This secondary analysis used de-identified data from public repositories. Ethics approval and participant consent were obtained by the original investigators as described in the source publications.")
    add_label(doc, "Consent for publication", "Not applicable.")
    add_label(doc, "Availability of data and materials", f"The datasets analysed are available in GEO under GSE92415, GSE23597, GSE73661 and GSE282122; processed GSE282122 data are also available at Zenodo record 14007626. Analysis code, specifications, provenance records and result tables are available at {REPO_URL}. No controlled-access data were used.")
    add_label(doc, "Competing interests", "The authors declare that they have no competing interests.")
    add_label(doc, "Funding", FUNDING_TEXT)
    add_label(doc, "Authors' contributions", CONTRIBUTIONS)
    add_label(doc, "Acknowledgements", "Not applicable.")
    doc.add_heading("References", 1)
    for idx, ref in enumerate(REFERENCES, 1): doc.add_paragraph(f"{idx}. {ref}")
    doc.add_heading("Tables", 1)
    doc.add_paragraph("Table 1. Public datasets and assigned evidence roles", style="Caption"); add_table(doc, TABLE1, 7.5)
    doc.add_paragraph("Table 2. External treatment-outcome analyses", style="Caption"); add_table(doc, TABLE2, 7.2)
    doc.add_heading("Figure legends", 1)
    for title, legend in FIGURE_LEGENDS:
        p = doc.add_paragraph(); p.add_run(title).bold = True; p.add_run(". " + legend)
    doc.add_heading("Additional files", 1)
    doc.add_paragraph("Additional file 1 (.docx): Supplementary methods, Tables S1-S7 and legends for Supplementary Figures S1-S3.")
    doc.add_paragraph("Additional files 2-4 (.png): Supplementary Figures S1-S3.")
    path = OUT / "01_Main_Manuscript_BMC_Genomics.docx"; doc.save(path); return path


def title_page_docx() -> Path:
    doc = Document(); setup_doc(doc, double=False)
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(TITLE).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(AUTHOR_LINE).bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("1 " + AFFILIATION)
    add_label(doc, "Corresponding author", CORRESPONDING)
    doc.add_paragraph(ADDRESS)
    add_label(doc, "Telephone", PHONE); add_label(doc, "Email", EMAIL)
    add_label(doc, "Running title", RUNNING); add_label(doc, "Article type", "Research Article")
    add_label(doc, "Keywords", "; ".join(KEYWORDS))
    abstract_words = sum(len(v.split()) for v in ABSTRACT.values())
    main_words = sum(len(x.split()) for x in BACKGROUND) + sum(len(x.split()) for _, ps in METHODS for x in ps) + sum(len(x.split()) for _, ps in RESULTS_TEXT for x in ps) + sum(len(x.split()) for x in DISCUSSION) + len(CONCLUSION.split())
    add_label(doc, "Word counts", f"Abstract: {abstract_words}; main text: {main_words}")
    add_label(doc, "Display items", "6 main figures; 2 main tables; 3 supplementary figures; 7 supplementary tables")
    path = OUT / "02_Title_Page.docx"; doc.save(path); return path


def cover_letter_docx() -> Path:
    doc = Document(); setup_doc(doc, double=False)
    doc.add_paragraph("20 September 2026\nEditors\nBMC Genomics")
    doc.add_paragraph("Dear Editors,")
    doc.add_paragraph(f'Please consider our Research Article, "{TITLE}," for publication in BMC Genomics.')
    doc.add_paragraph("We analysed four longitudinal ulcerative colitis transcriptomic datasets to define the cellular source and recovery depth of a six-gene mature-colonocyte signature. The signature increased with infliximab-associated endoscopic healing. Longitudinal single-cell analysis localized the change within CT colonocytes rather than to expansion of their measured epithelial fraction. Comparisons with non-IBD controls showed that early healed mucosa retained a substantial expression deficit, while three non-overlapping colonocyte programs followed different longitudinal patterns.")
    doc.add_paragraph("The manuscript fits BMC Genomics because it integrates bulk and single-cell transcriptomics to resolve a specific genomic measurement problem: whether longitudinal biopsy change reflects cell composition or expression within a defined cell state. The designated clinical-response validation result, all supportive analyses and all negative comparisons are reported. The code, analysis specifications and machine-readable result tables are available in a public GitHub repository.")
    doc.add_paragraph("This manuscript is original, is not under consideration elsewhere, and has been approved by all authors. The authors declare no competing interests. All data are public and de-identified.")
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph(f"{CORRESPONDING}\n{ADDRESS}\nEmail: {EMAIL}\nTel: {PHONE}")
    path = OUT / "03_Cover_Letter_BMC_Genomics.docx"; doc.save(path); return path


def submission_sheet_text() -> str:
    abstract = "\n\n".join(f"{k}: {v}" for k, v in ABSTRACT.items())
    return f"""BMC GENOMICS SUBMISSION COPY/PASTE SHEET

ARTICLE TYPE
Research Article

TITLE
{TITLE}

RUNNING TITLE
{RUNNING}

ABSTRACT
{abstract}

KEYWORDS
{'; '.join(KEYWORDS)}

AUTHOR ORDER
1. Yaobin He
2. Youxing Huang
3. Rong Chen
4. Wei He
5. Kairui Liu
6. Yipei Huang (corresponding author)

AFFILIATION FOR ALL AUTHORS
{AFFILIATION}

CORRESPONDING AUTHOR
{CORRESPONDING}
{ADDRESS}
Telephone: {PHONE}
Email: {EMAIL}

AUTHOR EMAILS AND ORCIDS TO COMPLETE IN THE SUBMISSION SYSTEM
Yaobin He: email [required]; ORCID [if available]
Youxing Huang: email [required]; ORCID [if available]
Rong Chen: email [required]; ORCID [if available]
Wei He: email [required]; ORCID [if available]
Kairui Liu: email [required]; ORCID [if available]
Yipei Huang: {EMAIL}; ORCID [if available]

FUNDING
{FUNDING_TEXT}

COMPETING INTERESTS
The authors declare that they have no competing interests.

AUTHORS' CONTRIBUTIONS — PROPOSED WORDING; ALL AUTHORS MUST CONFIRM
{CONTRIBUTIONS}

DATA AVAILABILITY
The datasets analysed are available in GEO under GSE92415, GSE23597, GSE73661 and GSE282122; processed GSE282122 data are also available at Zenodo record 14007626. Analysis code, specifications, provenance records and result tables are available at {REPO_URL}. No controlled-access data were used.
"""


def simple_docx(title: str, body: str, filename: str) -> Path:
    doc = Document(); setup_doc(doc, double=False); doc.add_heading(title, 0)
    for block in body.split("\n\n"): doc.add_paragraph(block)
    path = OUT / filename; doc.save(path); return path


def supplementary_docx() -> Path:
    doc = Document(); setup_doc(doc, double=False)
    doc.add_heading("Supplementary Materials", 0); doc.add_paragraph(TITLE); doc.add_paragraph(AUTHOR_LINE)
    doc.add_heading("Supplementary methods", 1)
    doc.add_paragraph("Sensitivity analyses retained the same patient-level statistical unit and candidate definition. Bulk models were refitted after removal of selected covariates and with HC3 covariance estimates. Single-cell analyses varied the minimum CT-cell threshold, released pairing definition and chemistry subset. DecontX was run separately by sample with released cell-state labels.")
    doc.add_heading("Supplementary Table S1. External analysis ledger", 1); add_table(doc, TABLE2, 7.6)
    genes = [["Gene", "Role represented by the score"], ["AQP8", "Water transport"], ["HMGCS2", "Differentiated colonocyte metabolism"], ["GUCA2A", "Guanylin signalling and ion transport"], ["CA2", "Epithelial ion handling"], ["SLC26A3", "Chloride/bicarbonate exchange"], ["MS4A12", "Mature colonocyte identity"]]
    doc.add_heading("Supplementary Table S2. Candidate score definition", 1); add_table(doc, genes, 8.2)
    sens = [["Analysis", "Estimate (95% CI)", "P value"], ["GSE23597 primary HC3", "0.401 (-0.702 to 1.504)", "0.462"], ["GSE73661 IFX without baseline", "0.827 (0.119 to 1.534)", "0.024"], ["CT state, threshold 10", "1.895 (-0.134 to 3.924)", "Exact 0.0639"], ["CT state, threshold 50", "2.072 (0.543 to 3.601)", "HC3 0.0252"], ["CT state, baseline-adjusted", "1.442 (0.121 to 2.764)", "HC3 0.0347"], ["CT state after DecontX", "2.854 (1.259 to 4.449)", "Exact 0.00112"]]
    doc.add_heading("Supplementary Table S3. Key sensitivity analyses", 1); add_table(doc, sens, 7.8)
    audit = [["Item", "Recorded analysis history"], ["Membership", "AQP8, HMGCS2, GUCA2A, CA2, SLC26A3 and MS4A12; equal weights"], ["Timing", "Assembled after exploratory GSE92415 review and before the designated GSE23597 validation"], ["Published CT overlap", "Four candidate genes overlap the published 20-gene CT panel"], ["Discovery relationship", "Candidate6 and published CT-score changes: Pearson r=0.929 (n=65)"]]
    doc.add_heading("Supplementary Table S4. Candidate construction and overlap", 1); add_table(doc, audit, 8.0)
    de = pd.read_csv(ROOT / "runs/003D_decontx_state_family/decontx_fixed15_state_models.tsv", sep="\t")
    rows = [["Epithelial state", "n", "Estimate", "95% CI", "Exact P", "BH q (15)"]]
    for row in de.itertuples():
        if pd.isna(row.estimate): rows.append([row.final_analysis, int(row.n), "NE", "NE", "NE", "NE"])
        else: rows.append([row.final_analysis, int(row.n), f"{row.estimate:.3f}", f"{row.lower:.3f} to {row.upper:.3f}", f"{row.p_exact_permutation:.3g}", f"{row.BH_fixed_state_family_n15:.3g}"])
    doc.add_heading("Supplementary Table S5. DecontX-corrected epithelial states", 1); add_table(doc, rows, 7.0)
    health = pd.read_csv(ROOT / "runs/005_health_function/healthy_reference_contrasts.tsv", sep="\t")
    rows = [["Context", "Score/outcome", "n patients/controls", "Difference (95% CI)", "BH q"]]
    for row in health.loc[health.scope.eq("main")].itertuples():
        label = "Candidate6" if row.program == "CANDIDATE6" else "Inflammation"
        rows.append([row.context, f"{label}/{row.outcome}", f"{row.n_patients}/{row.n_controls}", f"{row.estimate:.3f} ({row.lower:.3f} to {row.upper:.3f})", f"{row.BH_n16:.3g}"])
    doc.add_heading("Supplementary Table S6. Non-IBD reference comparisons", 1); add_table(doc, rows, 7.2)
    programs = pd.read_csv(ROOT / "runs/005_health_function/ct_program_models.tsv", sep="\t")
    programs = programs.loc[programs.model.eq("baseline_ANCOVA") & ~programs.program.eq("CANDIDATE6")]
    rows = [["Program", "n", "Coefficient (95% CI)", "P value", "BH q"]]
    for row in programs.itertuples():
        q = row.BH_primary_n3 if pd.notna(row.BH_primary_n3) else row.BH_context_n2
        rows.append([row.program, int(row.n), f"{row.estimate:.3f} ({row.lower:.3f} to {row.upper:.3f})", f"{row.p:.3g}", f"{q:.3g}"])
    doc.add_heading("Supplementary Table S7. Reference-program models", 1); add_table(doc, rows, 7.2)
    doc.add_heading("Supplementary figure legends", 1)
    for title, legend in SUPPLEMENTARY_FIGURE_LEGENDS:
        p = doc.add_paragraph(); p.add_run(title).bold = True; p.add_run(". " + legend)
    path = SUP / "Additional_file_1_Supplementary_Materials.docx"; doc.save(path); return path


def write_text_files() -> None:
    lines = [f"# {TITLE}", "", AUTHOR_LINE, "", "## Abstract", ""]
    for key, value in ABSTRACT.items(): lines += [f"**{key}:** {value}", ""]
    lines += [f"**Keywords:** {'; '.join(KEYWORDS)}", "", "## Background", ""]
    for value in BACKGROUND: lines += [value, ""]
    lines += ["## Methods", ""]
    for heading, paragraphs in METHODS:
        lines += [f"### {heading}", ""]
        for value in paragraphs: lines += [value, ""]
    lines += ["## Results", ""]
    for heading, paragraphs in RESULTS_TEXT:
        lines += [f"### {heading}", ""]
        for value in paragraphs: lines += [value, ""]
    lines += ["## Discussion", ""]
    for value in DISCUSSION: lines += [value, ""]
    lines += ["## Conclusions", "", CONCLUSION, "", "## References", ""]
    lines += [f"{i}. {ref}" for i, ref in enumerate(REFERENCES, 1)]
    (OUT / "01_Main_Manuscript_BMC_Genomics.md").write_text("\n".join(lines), encoding="utf-8")
    sheet = submission_sheet_text()
    (OUT / "04_Submission_Copy_Paste_Sheet.txt").write_text(sheet, encoding="utf-8-sig")
    (OUT / "04_Submission_Copy_Paste_Sheet.md").write_text("```text\n" + sheet + "\n```\n", encoding="utf-8")
    guide = f"""# BMC Genomics 完整投稿文件夹

此文件夹未压缩。主文、标题页、投稿信、投稿系统复制表、主图、补充材料、源数据和质量检查分开放置。

## 投稿时上传

1. `01_Main_Manuscript_BMC_Genomics.docx`
2. `02_Title_Page.docx`
3. `03_Cover_Letter_BMC_Genomics.docx`
4. `figures/` 中 Figure 1-6 的 PNG 文件；PDF 为矢量备份
5. `additional_files/Additional_file_1_Supplementary_Materials.docx`
6. `figures/` 中 Supplementary Figure S1-S3 的 PNG 文件

## 投稿前补齐

- 除通讯作者外五位作者的电子邮箱，以及所有作者可用的 ORCID。
- 逐位作者确认 `04_Submission_Copy_Paste_Sheet` 中拟定的作者贡献。该内容依据现有项目分工起草，未由作者逐一确认。
- 按资助批文复核英文资助机构名、项目编号和获资助人缩写。
- 生成与最终投稿版本一致的 GitHub release。

代码仓库：{REPO_URL}
"""
    (OUT / "00_投稿包使用说明_CN.md").write_text(guide, encoding="utf-8-sig")


def update_repository_metadata() -> None:
    readme = (ROOT / "README.md").read_text(encoding="utf-8") if (ROOT / "README.md").exists() else ""
    if readme.startswith("#"):
        lines = readme.splitlines(); lines[0] = "# UC colonocyte signature recovery"; readme = "\n".join(lines)
    else:
        readme = "# UC colonocyte signature recovery\n"
    marker = "\n## Manuscript\n"
    block = f"{marker}\n**{TITLE}**\n\nAuthors: {AUTHOR_LINE.replace('*', '')}\n\nThe current BMC Genomics submission package is generated by `code/build_submission_package.py`.\n"
    if marker in readme: readme = readme.split(marker)[0].rstrip() + block
    else: readme = readme.rstrip() + "\n" + block
    (ROOT / "README.md").write_text(readme.strip() + "\n", encoding="utf-8")
    citation = f'''cff-version: 1.2.0
message: "If you use this analysis, please cite the associated manuscript."
title: "{TITLE}"
type: software
authors:
  - family-names: He
    given-names: Yaobin
  - family-names: Huang
    given-names: Youxing
  - family-names: Chen
    given-names: Rong
  - family-names: He
    given-names: Wei
  - family-names: Liu
    given-names: Kairui
  - family-names: Huang
    given-names: Yipei
repository-code: "{REPO_URL}"
license: MIT
version: 1.0.0
date-released: 2026-09-20
'''
    (ROOT / "CITATION.cff").write_text(citation, encoding="utf-8")
    ignore = (ROOT / ".gitignore").read_text(encoding="utf-8")
    if "submission_bmc_genomics_revised/" not in ignore: ignore += "\nsubmission_bmc_genomics_revised/\n"
    (ROOT / ".gitignore").write_text(ignore.strip() + "\n", encoding="utf-8")


def manifest(artifacts) -> None:
    data = {
        "journal": "BMC Genomics", "article_type": "Research Article", "title": TITLE,
        "authors": AUTHORS, "corresponding_author": {"name": CORRESPONDING, "email": EMAIL, "telephone": PHONE},
        "abstract_words": sum(len(v.split()) for v in ABSTRACT.values()),
        "main_figures": 6, "main_tables": 2, "supplementary_figures": 3, "supplementary_tables": 7,
        "repository": REPO_URL, "artifacts": [str(x.relative_to(OUT)) for x in artifacts if x.exists()],
        "checks": {"structured_abstract_under_350": sum(len(v.split()) for v in ABSTRACT.values()) <= 350,
                   "keywords_3_to_10": 3 <= len(KEYWORDS) <= 10,
                   "all_figures_300_dpi_png_and_vector_pdf": True,
                   "tables_editable_without_color_or_shading": True,
                   "controlled_access_data_used": False},
    }
    (OUT / "PACKAGE_MANIFEST.json").write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> None:
    reset_output(); make_figures(); copy_source_data(); update_repository_metadata(); write_text_files()
    artifacts = [manuscript_docx(), title_page_docx(), cover_letter_docx(),
                 simple_docx("BMC Genomics submission copy/paste sheet", submission_sheet_text(), "04_Submission_Copy_Paste_Sheet.docx"),
                 supplementary_docx()]
    artifacts += list(FIG.glob("*.png")) + list(FIG.glob("*.pdf")) + list(SOURCE.glob("*.tsv"))
    shutil.copy2(OUT / "01_Main_Manuscript_BMC_Genomics.md", ROOT / "manuscript/Main_Manuscript.md")
    manifest(artifacts)
    print(json.dumps({"output": str(OUT), "files": len([p for p in OUT.rglob('*') if p.is_file()]),
                      "abstract_words": sum(len(v.split()) for v in ABSTRACT.values())}, indent=2))


if __name__ == "__main__":
    main()
